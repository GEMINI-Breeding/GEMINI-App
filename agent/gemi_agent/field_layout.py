"""Field layouts from the files breeders actually keep.

A field design arrives in many shapes; this reads them into one list of plots
(plot, tier, bed, genotype, crop, trial, block, rep) so the agent can reason
about where each plot is and which crop it is, and so results from an upload
that covers several crops can be filed per crop.

Shapes understood (any sheet of an .xlsx, or a .csv):

- **A table, one row per plot**, its header anywhere in the first rows, with
  the usual names in any spelling: `Field.plot.number`, `Field plot #`,
  `Plot`, `PlotNo`, `plots`; `Bed`/`Tier` (or `Row` for the tier), or one
  combined column ("Row 10-bed 64", "3 & 4"); the genotype as `genotype`,
  `entry.list.vect`, `Entry`, `Line ID`, `Variety Name`, `Codes`; the crop as
  `Crop`, `Population_exact`, `Population_general`, `Species`.
- **A drawn map**: tiers down the side, beds across ("Tier/bed #" then bed
  numbers), plot numbers in the cells ("B" = border); or beds down the side
  with a (plot, variety) pair of columns per tier ("Bed | Tier 1 | | Tier 2").
  A plot spanning two bed columns (30-inch beds) is one plot. A label row
  above the map ("COWPEA MAGIC POPULATION", "SORGHUM", merged across the beds
  it covers) gives each column's trial and crop.
- **Several of these at once** (sheets of one workbook, or several files): the
  plots are joined by plot number, the map giving positions and the table the
  genotype and crop.
"""
from __future__ import annotations

import csv
import itertools
import re
from pathlib import Path
from typing import Any, Optional

# ── names ───────────────────────────────────────────────────────────────────


def _norm(h: Any) -> str:
    return re.sub(r"[\s._#/()-]+", " ", str(h or "").strip().lower()).strip()


# In order of preference within each kind.
ALIASES: dict[str, tuple[str, ...]] = {
    "plot": ("field plot number", "plotid field", "plot id field", "field plot", "field plots", "plot number", "plotno", "plot no", "plot id",
             "plotid", "plot", "plots"),
    "tier": ("tier", "row", "range", "tier number", "row number"),
    "bed": ("bed", "col", "column", "bed number", "column number"),
    "rowbed": ("row bed number", "row bed", "bed tier"),
    "genotype": ("genotype", "entry list vect", "variety name", "accession", "line id", "line", "entry", "codes",
                 "variety", "germplasm", "name", "varietyname", "pedigree"),
    "crop": ("crop", "population exact", "population general", "population", "species"),
    "crop_hint": ("experiment", "locyearcropplot", "trial", "study"),
    "year": ("year", "season"),
    "location": ("location", "site", "loc", "environment", "env"),
    "block": ("block",),
    "rep": ("rep", "replicate", "replication"),
}
BORDER = re.compile(r"^(b|border|buffer|f|fill|filler|bf)$", re.I)

# Words in trial labels and crop columns → the crop.
CROP_WORDS = [
    (re.compile(r"\bcowpea|\bcp\b|vigna unguiculata|magic", re.I), "Cowpea"),
    (re.compile(r"tepary|acutifolius|interspecific", re.I), "Common bean / tepary"),
    (re.compile(r"common bean|\bcb\b|cdbn|phaseolus vulgaris|\bbean", re.I), "Common bean"),
    (re.compile(r"lima|lunatus", re.I), "Lima bean"),
    (re.compile(r"sorghum|\bs\b", re.I), "Sorghum"),
    (re.compile(r"garbanzo|chickpea", re.I), "Chickpea"),
]


def crop_of(label: Any) -> Optional[str]:
    t = str(label or "").strip()
    if not t:
        return None
    for pat, crop in CROP_WORDS:
        if pat.search(t):
            return crop
    return None


def _num(v: Any) -> Optional[float]:
    try:
        f = float(str(v).strip())
        return f
    except (TypeError, ValueError):
        return None


def _int(v: Any) -> Optional[int]:
    f = _num(v)
    return int(f) if f is not None and f == int(f) else None


def _plot_label(v: Any) -> Optional[str]:
    """A plot label as text: 1901, "KW7"; None for blanks and borders."""
    if v is None:
        return None
    t = str(v).strip()
    if not t or BORDER.match(t):
        return None
    f = _num(t)
    return str(int(f)) if f is not None and f == int(f) else t


# ── reading files ───────────────────────────────────────────────────────────


def _sheets(path: Path, sheet: Optional[str]) -> list[dict]:
    """Each sheet as {name, rows (strings), merged [(r0, c0, r1, c1)]} (0-based)."""
    if path.suffix.lower() in (".csv", ".txt", ".tsv"):
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            text = f.read()
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:                       # long rows of mostly empty cells: commas
            dialect = csv.excel
        rows = [[c.strip() for c in r] for r in csv.reader(text.splitlines(), dialect)]
        return [{"name": path.stem, "rows": rows, "merged": []}]
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("reading .xlsx needs openpyxl (pip install openpyxl), or save the sheet as .csv") from None
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        if sheet and ws.title != sheet:
            continue
        rows = [["" if c is None else str(c).strip() for c in r]
                for r in ws.iter_rows(max_row=min(ws.max_row, 5000), values_only=True)]
        while rows and not any(rows[-1]):
            rows.pop()
        merged = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
        out.append({"name": ws.title, "rows": rows, "merged": merged})
    if sheet and not out:
        raise ValueError(f"no sheet named {sheet!r}; sheets: {', '.join(w.title for w in wb.worksheets)}")
    return out


# ── one row per plot ────────────────────────────────────────────────────────


def _environment(year: str, location: str) -> dict:
    """Year and site of a row; "Kearney 2023" in one column gives both."""
    y = re.search(r"(19|20)\d\d", str(year or "")) or re.search(r"(19|20)\d\d", str(location or ""))
    loc = re.sub(r"\b(19|20)\d\d\b", "", str(location or "")).strip(" ,-_")
    return {"year": y.group(0) if y else "", "location": loc}


def _header_map(header: list[str]) -> dict[str, list[int]]:
    """kind -> candidate column indexes, best first."""
    names = [_norm(h) for h in header]
    found: dict[str, list[int]] = {}
    for kind, aliases in ALIASES.items():
        for a in aliases:
            for i, n in enumerate(names):
                if n == a and i not in found.get(kind, []):
                    found.setdefault(kind, []).append(i)
    # "row bed number" must not also count as a tier or bed column
    for i in found.get("rowbed", []):
        for k in ("tier", "bed"):
            if i in found.get(k, []):
                found[k].remove(i)
    return {k: v for k, v in found.items() if v}


def _long_table(sheet: dict) -> Optional[dict]:
    rows = sheet["rows"]
    for h in range(min(40, len(rows))):
        cols = _header_map(rows[h])
        has_pos = ("tier" in cols and "bed" in cols) or "rowbed" in cols
        if not ("plot" in cols and (has_pos or "genotype" in cols or "crop" in cols)) and not has_pos:
            continue
        body, blanks = [], 0
        for r in rows[h + 1:]:
            if not any(r):
                blanks += 1
                if blanks > 5:
                    break
                continue
            blanks = 0
            body.append(r + [""] * (len(rows[h]) - len(r)))
        if not body:
            continue

        def pick(kind: str, prefer_text: bool = False) -> Optional[int]:
            idx = cols.get(kind) or []
            if prefer_text:
                for i in idx:
                    vals = [r[i] for r in body[:200] if r[i]]
                    if vals and sum(_num(v) is None for v in vals) > len(vals) / 2:
                        return i
            return idx[0] if idx else None

        ip, it, ib, irb = pick("plot"), pick("tier"), pick("bed"), pick("rowbed")
        ig, ic = pick("genotype", prefer_text=True), pick("crop", prefer_text=True)
        ih = pick("crop_hint", prefer_text=True)
        iblock, irep = pick("block"), pick("rep")
        iyear, iloc = pick("year"), pick("location")
        plots = []
        for r in body:
            tier = _int(r[it]) if it is not None else None
            bed = span = None
            if ib is not None:
                nums = [int(x) for x in re.findall(r"\d+", r[ib])]
                if nums:
                    bed, span = nums[0], (nums[-1] - nums[0] + 1 if len(nums) > 1 else 1)
            if irb is not None and (tier is None or bed is None):
                m = re.search(r"row\s*(\d+).*?bed\s*(\d+)", r[irb], re.I) or re.search(r"(\d+)\D+(\d+)", r[irb])
                if m:
                    tier, bed = tier or int(m.group(1)), bed or int(m.group(2))
            plot = _plot_label(r[ip]) if ip is not None else None
            if plot is None and (tier is None or bed is None):
                continue
            plots.append({"columns": {rows[h][i]: r[i] for i in range(len(rows[h])) if rows[h][i] and r[i]},
                          "plot": plot, "tier": tier, "bed": bed, "bed_span": span or 1,
                          "genotype": r[ig] if ig is not None else "", "crop_label": r[ic] if ic is not None else "",
                          "crop_hint": r[ih] if ih is not None else "",
                          "block": r[iblock] if iblock is not None else "", "rep": r[irep] if irep is not None else "",
                          **_environment(r[iyear] if iyear is not None else "", r[iloc] if iloc is not None else "")})
        if plots:
            used = {k: rows[h][v] for k, v in (("plot", ip), ("tier", it), ("bed", ib), ("row+bed", irb),
                                                 ("genotype", ig), ("crop", ic), ("block", iblock), ("rep", irep))
                    if v is not None}
            return {"kind": "table", "sheet": sheet["name"], "header_row": h + 1, "columns_used": used,
                    "plots": plots}
    return None


# ── drawn maps ──────────────────────────────────────────────────────────────

_TIER_N = re.compile(r"^tier\s*(\d+)$", re.I)


def _labels_above(sheet: dict, header_row: int, col_from: int, ncols: int) -> dict[int, str]:
    """Column -> the trial label written above the map (merged across its columns).
    Of the few rows above the header, the one whose labels name crops or
    trials best (a row of sizes or pass numbers is not it)."""
    rows, merged = sheet["rows"], sheet["merged"]
    best: tuple[int, int, dict[int, str]] = (-1, -1, {})
    for r in range(header_row - 1, max(-1, header_row - 5), -1):
        row = rows[r] if r < len(rows) else []
        filled = [row[c] for c in range(col_from, min(len(row), ncols)) if row[c]]
        cells = {c: row[c] for c in range(col_from, min(len(row), ncols)) if row[c] and _num(row[c]) is None}
        if not cells or len(cells) < len(filled) / 2:
            continue
        lab: dict[int, str] = {}
        row_merged = [m for m in merged if m[0] <= r <= m[2]]
        for c, text in cells.items():
            span = next(((m[1], m[3]) for m in row_merged if m[1] <= c <= m[3]), (c, c))
            for k in range(span[0], span[1] + 1):
                lab[k] = text
        if not row_merged:                      # unmerged labels run until the next one
            last = None
            for k in range(col_from, ncols):
                if k in lab:
                    last = lab[k]
                elif last is not None:
                    lab[k] = last
        score = (sum(1 for t in cells.values() if crop_of(t)), len(cells))
        if score > best[:2]:
            best = (*score, lab)
    return best[2]


def _grid(sheet: dict) -> Optional[dict]:
    rows = sheet["rows"]
    for h in range(min(60, len(rows))):
        row = rows[h]
        # (a) "Tier/bed #" then bed numbers across; tiers down the side
        for c, cell in enumerate(row):
            n = _norm(cell)
            if ("bed" in n and ("tier" in n or n.startswith("row"))) and "dim" not in n and len(n) < 20:
                beds = {k: _int(row[k]) for k in range(c + 1, len(row)) if _int(row[k]) is not None}
                if len(beds) < 3:
                    continue
                labels = _labels_above(sheet, h, c + 1, len(row))
                cells: dict[str, dict] = {}
                for r in rows[h + 1:]:
                    tier = _int(r[c]) if c < len(r) else None
                    if tier is None:
                        continue
                    for k, bed in beds.items():
                        if k >= len(r):
                            continue
                        plot = _plot_label(r[k])
                        if plot is None or _int(plot) is None:
                            continue
                        p = cells.get(plot)
                        if p and p["tier"] == tier and bed == p["bed"] + p["bed_span"]:
                            p["bed_span"] += 1                    # a plot over two narrow beds
                            continue
                        if p:
                            continue                               # the same number elsewhere: keep the first
                        cells[plot] = {"plot": plot, "tier": tier, "bed": bed, "bed_span": 1,
                                       "trial": labels.get(k, ""), "genotype": "", "crop_label": ""}
                if cells:
                    return {"kind": "map (tiers down, beds across)", "sheet": sheet["name"], "header_row": h + 1,
                            "plots": list(cells.values()),
                            "trials": sorted({v for v in labels.values() if v})}
        # (b) "Bed | Tier 1 | | Tier 2 …": beds down, a (plot, variety) pair per tier
        tiers = {k: int(m.group(1)) for k, cell in enumerate(row) if (m := _TIER_N.match(str(cell).strip()))}
        if len(tiers) >= 2 and _norm(row[0]) in ("bed", "beds", "row", "rows", "bed number"):
            cells = {}
            heads = sorted(tiers)
            for r in rows[h + 1:]:
                bed = _int(r[0]) if r else None
                if bed is None:
                    continue
                # (plot, variety) pairs; each belongs to the nearest "Tier n" at or left of it
                # (the pairs are often drawn one column right of their header)
                for k in range(1, len(r)):
                    plot = _plot_label(r[k])
                    if plot is None or _int(plot) is None:
                        continue
                    at = [t for t in heads if t <= k]
                    if not at:
                        continue
                    nxt = r[k + 1] if k + 1 < len(r) else ""
                    if plot not in cells:
                        cells[plot] = {"plot": plot, "tier": tiers[at[-1]], "bed": bed, "bed_span": 1, "trial": "",
                                       "genotype": nxt if nxt and _num(nxt) is None else "", "crop_label": ""}
            if cells:
                notes = " ".join(" ".join(x for x in rr if x) for rr in rows[:h])
                return {"kind": "map (beds down, tiers across)", "sheet": sheet["name"], "header_row": h + 1,
                        "plots": list(cells.values()), "trials": [], "notes": notes[:300]}
    return None


# ── putting it together ─────────────────────────────────────────────────────


def read_layout(paths: list[str] | str, sheet: Optional[str] = None, year: Optional[str] = None,
                location: Optional[str] = None, experiment: Optional[str] = None) -> dict:
    """Every plot the files describe, with how each file was read.

    Tables with a year and site keep plots of different seasons and sites
    apart (plot numbers repeat across them); `year` / `location` keep one.
    A drawn map (which names neither) gives positions and crops to the plots
    with its plot numbers. Where two files place the same plot differently
    the conflict is listed (and a consistent bed/tier swap named)."""
    if isinstance(paths, str):
        paths = [paths]
    parts: list[dict] = []
    for p in paths:
        path = Path(p).expanduser()
        if not path.exists():
            raise ValueError(f"file not found: {p}")
        for sh in _sheets(path, sheet):
            got = _grid(sh) or _long_table(sh)
            if got:
                got["file"] = path.name
                parts.append(got)
    if not parts:
        raise ValueError("no plots found: no sheet has a recognisable plot table (plot / bed / tier columns) or "
                         "field map (a 'Tier/bed #' or 'Bed | Tier 1 …' header)")

    def related(a: Any, b: Any) -> bool:
        x, y = _norm(a), _norm(b)
        return bool(x and y and (x in y or y in x))

    def keep(p: dict) -> bool:
        if year and p.get("year") and str(year) not in str(p["year"]):
            return False
        if location and p.get("location") and not related(location, p["location"]):
            return False
        # A master sheet of several trials: its experiment / trial column.
        named = [p.get(k) for k in ("crop_hint", "trial") if p.get(k)]
        if experiment and named and not any(related(experiment, n) for n in named):
            return False
        return True

    plots: dict[tuple, dict] = {}
    unnumbered: list[dict] = []
    conflicts: list[dict] = []
    # Where files disagree on a plot's position, a field design (by its name) wins over other tables
    # (lab sheets, master tables), which come after it.
    tables = sorted((x for x in parts if x["kind"] == "table"),
                    key=lambda x: 0 if re.search(r"design|layout|field ?map|plot ?map", x["file"], re.I) else 1)
    maps = [x for x in parts if x["kind"] != "table"]
    sheets_per_file: dict[str, int] = {}
    for part in tables:
        sheets_per_file[part["file"]] = sheets_per_file.get(part["file"], 0) + 1
    for part in tables:
        # a single-crop file often says its crop only in its name ("..._Sorghum_R2.xlsx")
        by_name = crop_of(Path(part["file"]).stem.replace("_", " "))
        for p in part["plots"]:
            # a workbook with a sheet per group ("MAGIC", "CB-tepary", "Check")
            if not p.get("crop_hint") and crop_of(part["sheet"]):
                p["crop_hint"] = part["sheet"]
            if not p.get("crop_label") and sheets_per_file[part["file"]] > 1 and not crop_of(part["sheet"]):
                p["crop_label"] = part["sheet"]
            if not p.get("crop_hint") and by_name:
                p["crop_hint"] = by_name
            if not keep(p):
                continue
            if p["plot"] is None:
                unnumbered.append(p)
                continue
            key = (p.get("year") or "", (p.get("location") or "").lower(), p["plot"])
            cur = plots.get(key)
            if cur is None:
                plots[key] = {**p, "_from": part["file"]}
                continue
            _merge(cur, p, part["file"], conflicts)
    # maps: positions and crop labels for every plot with that number (or new plots)
    for part in maps:
        by_number: dict[str, list[dict]] = {}
        for cur in plots.values():
            by_number.setdefault(cur["plot"], []).append(cur)
        for p in part["plots"]:
            targets = by_number.get(p["plot"])
            if not targets:
                plots[("", "", p["plot"])] = {**p, "_from": part["file"]}
                continue
            for cur in targets:
                _merge(cur, p, part["file"], conflicts)
    rows = list(plots.values()) + unnumbered
    # A map-only file whose name names one crop ("..._Sorghum_map.xlsx"): its plots without a label.
    map_named = {part["file"]: crop_of(Path(part["file"]).stem.replace("_", " ")) for part in maps}
    for p in rows:
        p["crop"] = (crop_of(p.get("crop_label")) or p.get("crop_label") or crop_of(p.get("trial"))
                     or crop_of(p.get("crop_hint")) or map_named.get(p.get("_from")) or "")
        p["border"] = bool(re.search(r"(^|\s)(b|border)$", str(p.get("trial") or ""), re.I))
    swapped = [c for c in conflicts if c.get("swapped")]
    return {"plots": rows, "conflicts": conflicts,
            "swap_suspected": (len(swapped) > 0.8 * len(conflicts) and len(swapped) >= 3),
            "parts": [{k: v for k, v in part.items() if k != "plots"} | {"plots": len(part["plots"])}
                      for part in parts]}


def read_for_field(paths: list[str] | str, year: Optional[str] = None, location: Optional[str] = None,
                   experiment: Optional[str] = None) -> tuple[dict, dict, list[str]]:
    """The plots of one field from a file that may hold many (a master sheet
    of years, sites and trials). A filter that would keep nothing is dropped.
    Returns (layout, the whole file's layout, filters dropped)."""
    whole = read_layout(paths)
    want = {"year": year, "location": location, "experiment": experiment}
    if not any(want.values()):
        return whole, whole, []
    layout = read_layout(paths, **want)
    if layout["plots"]:
        return layout, whole, []
    # Drop as few filters as will find the field, the experiment first.
    given = [k for k in ("experiment", "location", "year") if want[k]]
    for n in range(1, len(given) + 1):
        for drop in itertools.combinations(given, n):
            layout = read_layout(paths, **{k: (None if k in drop else v) for k, v in want.items()})
            if layout["plots"]:
                return layout, whole, list(drop)
    return whole, whole, given


def _merge(cur: dict, p: dict, source: str, conflicts: list[dict]) -> None:
    """Fill `cur` from `p`; a different position for the same plot is a conflict."""
    if None not in (cur.get("tier"), cur.get("bed"), p.get("tier"), p.get("bed")) and \
            (cur["tier"], cur["bed"]) != (p["tier"], p["bed"]):
        conflicts.append({"plot": cur["plot"], "year": cur.get("year") or "", "location": cur.get("location") or "",
                          cur.get("_from", "?"): f"tier {cur['tier']}, bed {cur['bed']}",
                          source: f"tier {p['tier']}, bed {p['bed']}",
                          "swapped": (cur["tier"], cur["bed"]) == (p["bed"], p["tier"])})
    for k, v in p.items():
        if k == "columns":
            cur["columns"] = {**(v or {}), **(cur.get("columns") or {})}
        elif v not in (None, "") and cur.get(k) in (None, ""):
            cur[k] = v


def _ranges(nums: list[int]) -> str:
    nums = sorted(set(nums))
    if not nums:
        return ""
    out, a, b = [], nums[0], nums[0]
    for n in nums[1:]:
        if n == b + 1:
            b = n
        else:
            out.append(f"{a}-{b}" if a != b else f"{a}")
            a = b = n
    out.append(f"{a}-{b}" if a != b else f"{a}")
    return ", ".join(out)


def summarize(layout: dict) -> dict:
    plots = [p for p in layout["plots"] if not p.get("border")]
    by_crop: dict[str, list[dict]] = {}
    for p in plots:
        by_crop.setdefault(p.get("crop") or "(no crop given)", []).append(p)
    crops = []
    for crop, ps in sorted(by_crop.items(), key=lambda kv: -len(kv[1])):
        crops.append({"crop": crop, "plots": len(ps),
                      "beds": _ranges([p["bed"] for p in ps if p.get("bed") is not None]),
                      "tiers": _ranges([p["tier"] for p in ps if p.get("tier") is not None]),
                      "trials": sorted({p.get("trial") for p in ps if p.get("trial")})[:6],
                      "genotypes": len({p.get("genotype") for p in ps if p.get("genotype")}),
                      "plot_numbers": _ranges([int(p["plot"]) for p in ps if _int(p.get("plot")) is not None])})
    missing = [p["plot"] for p in plots if p.get("bed") is None or p.get("tier") is None]
    envs: dict[str, int] = {}
    for p in plots:
        e = " ".join(x for x in (p.get("location"), p.get("year")) if x)
        if e:
            envs[e] = envs.get(e, 0) + 1
    conflicts = layout.get("conflicts") or []
    return {"plots": len(plots), "crops": crops, "read_as": layout["parts"],
            "environments": envs or None,
            "position_conflicts": len(conflicts) or None,
            "conflict_examples": conflicts[:4] or None,
            "bed_tier_swapped": ("the files disagree on bed and tier for the same plots in a consistent way: one of "
                                 "them has the two columns the other way round. Ask the person which is right."
                                 if layout.get("swap_suspected") else None),
            "without_position": len(missing), "examples_without_position": missing[:8],
            "border_plots": sum(1 for p in layout["plots"] if p.get("border"))}


def write_csv(layout: dict, out: Path) -> Path:
    """The plots as the app's field-design CSV (row = tier, col = bed)."""
    out.parent.mkdir(parents=True, exist_ok=True)
    cols = ["plot", "row", "col", "bed_span", "accession", "crop", "trial", "block", "rep"]
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for p in layout["plots"]:
            if p.get("border"):
                continue
            w.writerow([p.get("plot") or "", p.get("tier") if p.get("tier") is not None else "",
                        p.get("bed") if p.get("bed") is not None else "", p.get("bed_span") or 1,
                        p.get("genotype") or "", p.get("crop") or "", p.get("trial") or "",
                        p.get("block") or "", p.get("rep") or ""])
    return out
