"""Plot boundaries without drawing: find the plots on an orthomosaic, label
them from a field design file, save them as the app's plot-boundary version,
and cut the orthomosaic into plot images.

This is the agent's way through the Process tab's "Plot Boundary Prep" and
"Split Into Plot Images" steps. The geometry is found by measuring
vegetation on the orthomosaic (beds across, tiers along the field); the
agent's part is to give the layout facts, look at the overlay it gets back,
and correct what is off.
"""
from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path
from typing import Any, Optional

from .client import AGENT_DIR
from .process_runs import scope_from_path
from .tools import _bool, _int, _str, _tool, client, process_tab

WORK = AGENT_DIR / "work" / "ortho"
_num = {"type": "number"}


# ── georeferencing (GeoTIFF tags + UTM, no GIS libraries) ────────────────────

def _utm_to_lonlat(e: float, n: float, zone: int, north: bool) -> tuple[float, float]:
    a, f = 6378137.0, 1 / 298.257223563
    k0, e2 = 0.9996, f * (2 - f)
    ep2 = e2 / (1 - e2)
    x, y = e - 500000.0, n if north else n - 10000000.0
    m = y / k0
    mu = m / (a * (1 - e2 / 4 - 3 * e2 ** 2 / 64 - 5 * e2 ** 3 / 256))
    e1 = (1 - math.sqrt(1 - e2)) / (1 + math.sqrt(1 - e2))
    phi1 = (mu + (3 * e1 / 2 - 27 * e1 ** 3 / 32) * math.sin(2 * mu) + (21 * e1 ** 2 / 16 - 55 * e1 ** 4 / 32) * math.sin(4 * mu)
            + (151 * e1 ** 3 / 96) * math.sin(6 * mu) + (1097 * e1 ** 4 / 512) * math.sin(8 * mu))
    c1, t1 = ep2 * math.cos(phi1) ** 2, math.tan(phi1) ** 2
    n1 = a / math.sqrt(1 - e2 * math.sin(phi1) ** 2)
    r1 = a * (1 - e2) / (1 - e2 * math.sin(phi1) ** 2) ** 1.5
    d = x / (n1 * k0)
    lat = phi1 - (n1 * math.tan(phi1) / r1) * (
        d ** 2 / 2 - (5 + 3 * t1 + 10 * c1 - 4 * c1 ** 2 - 9 * ep2) * d ** 4 / 24
        + (61 + 90 * t1 + 298 * c1 + 45 * t1 ** 2 - 252 * ep2 - 3 * c1 ** 2) * d ** 6 / 720)
    lon = (d - (1 + 2 * t1 + c1) * d ** 3 / 6
           + (5 - 2 * c1 + 28 * t1 - 3 * c1 ** 2 + 8 * ep2 + 24 * t1 ** 2) * d ** 5 / 120) / math.cos(phi1)
    return math.degrees(lon) + (zone - 1) * 6 - 180 + 3, math.degrees(lat)


class Ortho:
    """An orthomosaic from storage: pixels (reduced if large) and the mapping
    from pixel to longitude/latitude."""

    def __init__(self, path: str, max_side: int = 5000):
        from PIL import Image

        Image.MAX_IMAGE_PIXELS = None
        self.path = path
        WORK.mkdir(parents=True, exist_ok=True)
        local = WORK / (re.sub(r"[^A-Za-z0-9._-]", "_", path)[-150:])
        if not local.exists():
            local.write_bytes(client().download(path))
        im = Image.open(local)
        tags = im.tag_v2
        scale, tie, keys = tags.get(33550), tags.get(33922), tags.get(34735)
        if not scale or not tie:
            raise ValueError("this file has no georeferencing (not a GeoTIFF)")
        epsg = None
        if keys:
            for i in range(4, len(keys), 4):
                if keys[i] == 3072:
                    epsg = int(keys[i + 3])
        self.epsg = epsg
        full_w, full_h = im.size
        frame = 0
        for i in range(getattr(im, "n_frames", 1)):      # smallest overview that is still detailed enough
            im.seek(i)
            if max(im.size) <= max_side:
                frame = i
                break
            frame = i
        im.seek(frame)
        self.image = im.convert("RGBA")
        self.w, self.h = self.image.size
        k = full_w / self.w
        self.px_x, self.px_y = float(scale[0]) * k, float(scale[1]) * k      # metres (or degrees) per pixel
        self.x0 = float(tie[3]) - float(tie[0]) * float(scale[0])
        self.y0 = float(tie[4]) + float(tie[1]) * float(scale[1])
        if epsg and 32601 <= epsg <= 32660:
            self.zone, self.north = epsg - 32600, True
        elif epsg and 32701 <= epsg <= 32760:
            self.zone, self.north = epsg - 32700, False
        elif epsg == 4326:
            self.zone = 0
            self.north = True
        else:
            raise ValueError(f"unsupported map projection (EPSG {epsg}); expected UTM or WGS84")

    @property
    def metres_per_px(self) -> float:
        return self.px_x if self.zone else self.px_x * 111_320 * math.cos(math.radians(self.y0))

    def lonlat(self, col: float, row: float) -> list[float]:
        x, y = self.x0 + col * self.px_x, self.y0 - row * self.px_y
        if not self.zone:
            return [round(x, 8), round(y, 8)]
        lon, lat = _utm_to_lonlat(x, y, self.zone, self.north)
        return [round(lon, 8), round(lat, 8)]


# ── finding the plots ────────────────────────────────────────────────────────

def _runs(mask) -> list[tuple[int, int]]:
    out, start = [], None
    for i, v in enumerate(list(mask) + [False]):
        if v and start is None:
            start = i
        elif not v and start is not None:
            out.append((start, i))
            start = None
    return out


def _smooth(a, k: int):
    import numpy as np

    k = max(1, int(k))
    return np.convolve(np.pad(a, k // 2, mode="edge"), np.ones(k) / k, mode="valid")[: len(a)]


def find_layout(ortho: Ortho, beds: int, tiers: int, bed_axis: str = "auto", border_tiers: Optional[int] = None,
                block: int = 0) -> dict:
    """Locate the planted block and its tier strips from vegetation.

    Tiers are the strips separated by alleys along the field's long axis;
    beds divide the block evenly across it. Returns pixel geometry and what
    was assumed, for the overlay and the caller's judgement.
    """
    import numpy as np

    a = np.asarray(ortho.image, dtype=np.float32)
    rgb, alpha = a[..., :3] / 255.0, a[..., 3] > 0
    exg = (2 * rgb[..., 1] - rgb[..., 0] - rgb[..., 2]) / (rgb.sum(axis=2) + 1e-6)
    # Threshold between soil and canopy chosen from this image (Otsu), so a
    # dark early-morning flight is handled like a bright midday one.
    vals = exg[alpha]
    hist, edges = np.histogram(vals, bins=128, range=(-0.2, 0.5))
    w = np.cumsum(hist).astype(np.float64)
    mu = np.cumsum(hist * (edges[:-1] + edges[1:]) / 2)
    between = (mu[-1] * w - mu * w[-1]) ** 2 / (w * (w[-1] - w) + 1e-9)
    thr = float(edges[int(np.argmax(between[:-1]))])
    veg = (exg > min(max(thr, 0.02), 0.12)) & alpha
    long_is_rows = ortho.h >= ortho.w if bed_axis == "auto" else bed_axis == "columns"
    v = veg if long_is_rows else veg.T                  # v[along, across]
    m = ortho.metres_per_px

    # Across the field: the planted blocks are the wide stretches of vegetation.
    across = _smooth(v.mean(axis=0), 0.3 / m)
    blocks = [r for r in _runs(across > 0.35 * across.max()) if (r[1] - r[0]) * m > 2.0]
    if not blocks:
        raise ValueError("no planted block found on this orthomosaic (is it mostly bare soil?)")
    blocks.sort(key=lambda r: -(r[1] - r[0]))
    if block >= len(blocks):
        raise ValueError(f"only {len(blocks)} block(s) found")
    u0, u1 = blocks[block]

    # Along the field, inside that block: strips of plants separated by alleys.
    along = _smooth(v[:, u0:u1].mean(axis=1), 0.25 / m)
    strips = [r for r in _runs(along > 0.30 * along.max()) if (r[1] - r[0]) * m > 0.8]
    # Where canopies have closed over an alley, several strips come out as
    # one long run: cut such runs into equal parts at the usual pitch.
    if len(strips) >= 4:
        starts = np.diff([a_ for a_, _ in strips])
        pitch = float(np.median(starts))
        usual = float(np.median([b - a_ for a_, b in strips]))
        split = []
        for a_, b in strips:
            n = int(round((b - a_ + (pitch - usual)) / pitch))
            if n <= 1:
                split.append((a_, b))
            else:
                step = (b - a_ + (pitch - usual)) / n
                split += [(int(a_ + i * step), int(a_ + i * step + usual)) for i in range(n)]
        strips = split
    found = len(strips)
    note = ""
    if border_tiers is None:
        border_tiers = (found - tiers) // 2 if found > tiers and (found - tiers) % 2 == 0 else 0
    if found - 2 * border_tiers == tiers:
        used = strips[border_tiers: found - border_tiers] if border_tiers else strips
        if border_tiers:
            note = f"{found} strips found; the first and last {border_tiers} taken as border, not plots"
    else:
        # Counts do not match the design: fall back to even spacing between the outermost strips.
        lo, hi = strips[0][0], strips[-1][1]
        step = (hi - lo) / tiers
        used = [(int(lo + i * step), int(lo + (i + 1) * step)) for i in range(tiers)]
        note = (f"WARNING: {found} strips found but the design has {tiers} tiers "
                f"(with border_tiers={border_tiers}); evenly spaced cells were used instead - check the overlay")
    bed_w = (u1 - u0) / beds
    return {"long_is_rows": long_is_rows, "block": [int(u0), int(u1)], "tiers": [list(map(int, s)) for s in used],
            "bed_width_px": bed_w, "strips_found": found, "blocks_found": len(blocks), "note": note,
            "block_width_m": round((u1 - u0) * m, 2), "field_length_m": round((used[-1][1] - used[0][0]) * m, 1),
            "tier_length_m": round(sum(b - a_ for a_, b in used) / len(used) * m, 2)}


# ── field design ─────────────────────────────────────────────────────────────

ALIASES = {"row": ("tier", "row", "range", "row_num"), "col": ("bed", "col", "column", "col_num"),
           "plot": ("field plot #", "plot", "plot_id", "plot_no", "plot_number", "plotid"),
           "accession": ("genotype", "accession", "entry", "variety", "line", "name")}


def read_design(path: str) -> dict:
    """A field design table (.xlsx or .csv): one row per plot with its tier
    (row), bed (column), plot label and genotype."""
    p = Path(path).expanduser()
    if not p.exists():
        raise ValueError(f"field design not found: {path}")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise RuntimeError("reading .xlsx needs openpyxl (pip install openpyxl), or save the sheet as .csv") from None
        ws = openpyxl.load_workbook(p, data_only=True).active
        table = [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(p, newline="", encoding="utf-8-sig") as f:
            table = [[c.strip() for c in r] for r in csv.reader(f)]
    header = [h.lower() for h in table[0]]
    cols = {}
    for key, names in ALIASES.items():
        cols[key] = next((header.index(n) for n in names if n in header), None)
    if cols["row"] is None or cols["col"] is None:
        raise ValueError(f"the design needs tier/row and bed/column columns; found {table[0]}")
    rows = []
    for r in table[1:]:
        if not any(r) or not r[cols["row"]] or not r[cols["col"]]:
            continue
        rec = {table[0][i]: r[i] for i in range(len(table[0])) if table[0][i]}
        rec["_row"], rec["_col"] = int(float(r[cols["row"]])), int(float(r[cols["col"]]))
        rec["_plot"] = r[cols["plot"]] if cols["plot"] is not None else f"{rec['_row']}-{rec['_col']}"
        rec["_accession"] = r[cols["accession"]] if cols["accession"] is not None else ""
        rows.append(rec)
    return {"rows": rows, "beds": sorted({r["_col"] for r in rows}), "tiers": sorted({r["_row"] for r in rows}),
            "columns": [h for h in table[0] if h]}


def field_design(args: dict) -> dict:
    d = read_design(args["path"])
    beds, tiers = d["beds"], d["tiers"]
    groups, start = [], beds[0]
    for a_, b in zip(beds, beds[1:] + [None]):
        if b is None or b != a_ + 1:
            groups.append([start, a_])
            start = b
    grid = {(r["_row"], r["_col"]): r["_plot"] for r in d["rows"]}
    corners = {f"tier {t}, bed {b}": grid.get((t, b)) for t in (tiers[0], tiers[-1]) for b in (beds[0], beds[-1])}
    return {"result": {"plots": len(d["rows"]), "beds": beds, "tiers": [tiers[0], tiers[-1]], "tier_count": len(tiers),
                       "bed_blocks": groups, "columns": d["columns"], "corner_plots": corners,
                       "genotypes": len({r["_accession"] for r in d["rows"]}),
                       "note": "bed_blocks = runs of consecutive bed numbers; a gap between runs is usually an "
                               "alley or a different treatment block, often flown and mapped separately"}}


# ── boundaries ───────────────────────────────────────────────────────────────

def numeric_plot(label: str, seen: dict) -> int:
    """The app stores plot numbers as integers. A design label that is not a
    number (check plots such as KW7, KD12) gets a stable numeric alias —
    9000 + 100 per distinct prefix + its own number: KW7 -> 9107, KD12 -> 9212 —
    and keeps its label in `plot_label`."""
    try:
        return int(float(label))
    except (TypeError, ValueError):
        pass
    m = re.match(r"^(.*?)(\d+)$", str(label).strip())
    prefix, n = (m.group(1), int(m.group(2))) if m else (str(label), 0)
    if prefix not in seen:
        seen[prefix] = 9000 + 100 * (len(seen) + 1)
    return seen[prefix] + n


def _cells(ortho: Ortho, lay: dict, design: dict, beds: list[int], tier_first_at: str, bed_first_at: str,
           shrink_m: float) -> list[dict]:
    """One polygon per (tier, bed), labelled from the design."""
    tiers = design["tiers"]
    by = {(r["_row"], r["_col"]): r for r in design["rows"]}
    long_rows = lay["long_is_rows"]
    u0 = lay["block"][0]
    bw = lay["bed_width_px"]
    pad = shrink_m / ortho.metres_per_px
    # Position order along the field (index 0 = top of image = north) and across it (0 = left = west).
    tier_order = tiers if tier_first_at in ("north", "west") else tiers[::-1]
    bed_order = beds if bed_first_at in ("west", "north") else beds[::-1]
    feats = []
    aliases: dict = {}
    for r in sorted(design["rows"], key=lambda r: (r["_col"], r["_row"])):   # same alias whichever flight
        numeric_plot(r["_plot"], aliases)
    for ti, (s0, s1) in enumerate(lay["tiers"]):
        for bi in range(len(beds)):
            a0, a1 = s0 + pad, s1 - pad
            c0, c1 = u0 + bi * bw + pad, u0 + (bi + 1) * bw - pad
            box = [(c0, a0), (c1, a0), (c1, a1), (c0, a1)] if long_rows else [(a0, c0), (a1, c0), (a1, c1), (a0, c1)]
            ring = [ortho.lonlat(x, y) for x, y in box]
            tier, bed = tier_order[ti], bed_order[bi]
            rec = by.get((tier, bed))
            label = rec["_plot"] if rec else f"{tier}-{bed}"
            props: dict[str, Any] = {"row": tier, "col": bed, "plot": numeric_plot(label, aliases) if rec else tier * 1000 + bed,
                                     "plot_label": label, "accession": rec["_accession"] if rec else "",
                                     "_px": [[round(x, 1), round(y, 1)] for x, y in box]}
            if rec:
                props.update({k: v for k, v in rec.items() if not k.startswith("_")})
            feats.append({"type": "Feature", "properties": props,
                          "geometry": {"type": "Polygon", "coordinates": [ring + [ring[0]]]}})
    return feats


def _overlay(ortho: Ortho, feats: list[dict], name: str, lay: dict) -> Path:
    from PIL import Image, ImageDraw, ImageFont

    img = ortho.image.convert("RGB")
    k = min(3.0, 2600 / max(img.size))
    if lay["long_is_rows"]:                    # lay the field on its side so it fits a screen
        img = img.rotate(90, expand=True)
    img = img.resize((round(img.width * k), round(img.height * k)))
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default(size=max(11, int(0.5 / ortho.metres_per_px * k)))
    except TypeError:
        font = ImageFont.load_default()

    def pt(x: float, y: float) -> tuple[float, float]:
        return (y * k, (ortho.w - x) * k) if lay["long_is_rows"] else (x * k, y * k)

    for f in feats:
        ring = [pt(x, y) for x, y in f["properties"]["_px"]]
        d.polygon(ring, outline=(255, 0, 255))
        cx, cy = sum(p[0] for p in ring) / 4, sum(p[1] for p in ring) / 4
        label = str(f["properties"].get("plot_label") or f["properties"]["plot"])
        box = d.textbbox((cx, cy), label, font=font, anchor="mm")
        d.rectangle(box, fill=(0, 0, 0))
        d.text((cx, cy), label, fill=(255, 255, 0), font=font, anchor="mm")
    compass = "north is LEFT, east is UP" if lay["long_is_rows"] else "north is UP, east is RIGHT"
    d.rectangle((0, 0, 330, 24), fill=(0, 0, 0))
    d.text((6, 4), compass, fill=(255, 255, 255))
    out = WORK / f"boundaries_{name}.png"
    img.save(out)
    return out


def plot_boundaries(args: dict) -> dict:
    """Find the plots on an orthomosaic, label them from the design, and
    (unless preview) save + activate them as the scope's plot boundaries."""
    path = args["orthomosaic_path"]
    scope = scope_from_path(path)
    if scope is None:
        raise ValueError("orthomosaic_path must be under Processed/{year}/{experiment}/{site}/{population}/{date}/…")
    design = read_design(args["design_path"])
    beds = [int(b) for b in args["beds"]] if args.get("beds") else design["beds"]
    ortho = Ortho(path)
    lay = find_layout(ortho, len(beds), len(design["tiers"]), border_tiers=args.get("border_tiers"),
                      block=int(args.get("block") or 0))
    feats = _cells(ortho, lay, design, beds, str(args.get("first_tier_at") or "south"),
                   str(args.get("first_bed_at") or "west"), float(args.get("shrink_m") or 0.0))
    name = re.sub(r"[^A-Za-z0-9]+", "_", "_".join(str(scope[k]) for k in ("location", "population", "date"))
                  + "_" + "-".join(map(str, (beds[0], beds[-1]))))
    image = _overlay(ortho, feats, name, lay)
    corners = {f"{ns}-{ew}": None for ns in ("north", "south") for ew in ("west", "east")}
    t, b = lay["tiers"], len(beds)
    idx = {"north-west": 0, "north-east": b - 1, "south-west": (len(t) - 1) * b, "south-east": len(t) * b - 1}
    if not lay["long_is_rows"]:
        idx = {"north-west": 0, "south-west": b - 1, "north-east": (len(t) - 1) * b, "south-east": len(t) * b - 1}
    for kname, i in idx.items():
        corners[kname] = feats[i]["properties"]["plot_label"]
    result: dict[str, Any] = {
        "plots": len(feats), "beds": beds, "tiers": len(t), "layout": {k: lay[k] for k in (
            "strips_found", "blocks_found", "block_width_m", "field_length_m", "tier_length_m", "note")},
        "bed_width_m": round(lay["bed_width_px"] * ortho.metres_per_px, 2),
        "plots_at_corners": corners,
        "assumed": {"first_tier_at": args.get("first_tier_at") or "south (default - confirm with the person)",
                    "first_bed_at": args.get("first_bed_at") or "west (default)"},
    }
    if args.get("preview"):
        result["saved"] = False
        return {"result": result, "images": [str(image)]}

    c = client()
    directory = "Processed/" + "/".join(str(scope[k]) for k in ("year", "experiment", "location", "population", "date",
                                                                "platform", "sensor")) + "/"
    existing: list[dict] = []
    if args.get("add_to_active"):
        try:
            snap = c.load_version(directory).get("state_snapshot") or {}
            keep = {(f["properties"]["row"], f["properties"]["col"]) for f in feats}
            existing = [f for f in (snap.get("boundaries") or {}).get("features") or []
                        if (f["properties"].get("row"), f["properties"].get("col")) not in keep]
        except Exception:
            existing = []
    clean = [{**f, "properties": {k: v for k, v in f["properties"].items() if k != "_px"}} for f in feats]
    snapshot = {"boundaries": {"type": "FeatureCollection", "features": existing + clean},
                "grid": {"rows": len(t), "cols": b, "spacing_m": 0, "angle_deg": 0, "row_offset": 0, "col_offset": 0,
                         "fill_pattern": "row-major"},
                "created_from": "grid",
                "field_design": {"source": Path(args["design_path"]).name, "by": "agent"}}
    v = c.save_version(directory, snapshot, args.get("name") or f"agent: {len(existing) + len(clean)} plots")
    c.post_json("/api/plot_geometry/versions/activate", {"directory": directory, "version": v["version"]})

    def update(pr):
        run = pr.ensure_run(scope)
        pr.set_step(run, "plot_boundary_prep", "completed",
                    outputs={"activeVersion": v["version"], "activeVersionName": v.get("name")})
        return pr.describe(run)

    result.update({"saved": True, "version": v["version"], "directory": directory,
                   "total_plots_in_version": len(existing) + len(clean), "process_tab": process_tab(update)})
    return {"result": result, "images": [str(image)]}


def split_plots(args: dict) -> dict:
    """SPLIT_ORTHOMOSAIC: cut an orthomosaic into one image per plot of the active boundaries."""
    path = args["orthomosaic_path"]
    scope = scope_from_path(path)
    if scope is None:
        raise ValueError("orthomosaic_path must be under Processed/…")
    c = client()
    directory = "Processed/" + "/".join(str(scope[k]) for k in ("year", "experiment", "location", "population", "date",
                                                                "platform", "sensor")) + "/"
    snap = c.load_version(directory).get("state_snapshot") or {}
    boundaries = snap.get("boundaries") or {}
    if not boundaries.get("features"):
        raise ValueError("no active plot boundaries for this scope: run plot_boundaries first")
    params = {**{k: scope[k] for k in ("year", "experiment", "location", "population", "date", "platform", "sensor")},
              "boundaries": boundaries, "orthomosaic_path": path}
    exp = next((e for e in c.entities("experiments") if e.get("experiment_name") == scope["experiment"]), None)
    job = c.submit_job("SPLIT_ORTHOMOSAIC", params, str(exp["id"]) if exp else None)

    def update(pr):
        return pr.record_job(scope, job)

    return {"result": {"job_id": job.get("id"), "status": job.get("status"), "plots": len(boundaries["features"]),
                       "output": directory + "PlotImages/", "process_tab": process_tab(update)}}


def boundaries_geojson(args: dict) -> dict:
    """The active plot boundaries of a scope (what run_inference needs as `boundaries`)."""
    scope = scope_from_path(args["path"])
    if scope is None:
        raise ValueError("path must be under Raw/ or Processed/ …")
    directory = "Processed/" + "/".join(str(scope[k]) for k in ("year", "experiment", "location", "population", "date",
                                                                "platform", "sensor")) + "/"
    v = client().load_version(directory)
    feats = ((v.get("state_snapshot") or {}).get("boundaries") or {}).get("features") or []
    return {"result": {"directory": directory, "version": v.get("version"), "plots": len(feats),
                       "sample": [f["properties"] for f in feats[:3]]}}


def _field(fn, description: str, props: dict, required: list[str]) -> None:
    _tool(fn, description, props, required, track=False, group="aerial")


_field(field_design, "Read a field design file (.xlsx / .csv) on this computer: plots, beds, tiers, bed blocks, "
       "which plot is at each corner of the design.", {"path": _str}, ["path"])
_field(plot_boundaries, "Find the plots on an orthomosaic from its vegetation (tier strips along the field, beds "
       "across), label them from the field design, and return an overlay picture. With preview=false it saves and "
       "activates them as the scope's plot boundaries (the Process tab's Plot Boundary Prep step).",
       {"orthomosaic_path": _str, "design_path": {**_str, "description": "field design file on this computer"},
        "beds": {"type": "array", "items": _int, "description": "the design's bed numbers that are on THIS "
                 "orthomosaic, west to east in the design's numbering (default: all beds)"},
        "first_tier_at": {**_str, "description": "south | north: which end of the field the lowest tier number is at"},
        "first_bed_at": {**_str, "description": "west | east: which side the lowest of these bed numbers is on"},
        "border_tiers": {**_int, "description": "unplanned strips at each end to ignore (default: worked out)"},
        "block": {**_int, "description": "0 = widest planted block (default), 1 = next widest"},
        "shrink_m": {**_num, "description": "shrink each plot by this margin on every side"},
        "add_to_active": {**_bool, "description": "keep the other plots of the active version (second flight of the same field)"},
        "preview": _bool, "name": _str}, ["orthomosaic_path", "design_path"])
_field(split_plots, "Cut an orthomosaic into one image per plot using the active plot boundaries "
       "(SPLIT_ORTHOMOSAIC; needs the geo worker).", {"orthomosaic_path": _str}, ["orthomosaic_path"])
_field(boundaries_geojson, "Which plot-boundary version is active for a scope, and how many plots it has.",
       {"path": _str}, ["path"])


# ── ground (rover) plots ─────────────────────────────────────────────────────

def mock_field_design(args: dict) -> dict:
    """A stand-in field design for a rover track that has none yet: bed =
    pass, tier = position within the bed counted from one end of the field,
    placeholder genotype names. For trying the pipeline; replace it with the
    real design when there is one."""
    from .ground import Markings
    from .tools import _track

    track = _track(args)
    m = Markings(track)
    if not m.plots:
        raise ValueError("no plot markings for this track yet")
    passes = track.passes()
    genotypes = int(args.get("genotypes") or 40)
    rows = []
    for p in passes:
        mine = [pl for pl in m.plots if pl["start"] is not None and p["first"] <= pl["start"] <= p["last"]]
        mine.sort(key=lambda pl: pl["start"])
        if p["heading"] in ("South", "West"):          # tier 1 at the south/west end whatever the driving direction
            mine = mine[::-1]
        for k, pl in enumerate(mine, 1):
            rows.append({"plot": pl["plot_id"], "bed": p["pass"], "tier": k,
                         "genotype": f"MOCK-{(pl['plot_id'] * 7 + p['pass'] * 3) % genotypes + 1:03d}"})
    out = Path(args.get("save_as") or (track.cache / "mock_field_design.csv"))
    with open(out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["plot", "bed", "tier", "genotype"])
        w.writeheader()
        w.writerows(rows)
    return {"result": {"file": str(out), "plots": len(rows), "beds": len({r["bed"] for r in rows}),
                       "tiers": max(r["tier"] for r in rows), "mock": True,
                       "note": "genotype names are placeholders; bed = rover pass, tier counted from the south end"}}


def ground_plot_boundaries(args: dict) -> dict:
    """Plot boundaries for a rover track from a stitch run: each stitched
    plot's georeferenced footprint, labelled from the field design. Saved and
    activated like the aerial ones (the ground run's Plot Boundary Prep)."""
    prefix = args["stitch_prefix"].rstrip("/") + "/"
    scope = scope_from_path(prefix)
    if scope is None:
        raise ValueError("stitch_prefix must be Processed/…/AgRowStitch_v{N}/")
    c = client()
    manifest = json.loads(c.download(prefix + "stitch_manifest.json"))
    design = read_design(args["design_path"])
    by_plot = {str(r["_plot"]): r for r in design["rows"]}
    feats, missing = [], 0
    for pid, info in (manifest.get("plots") or {}).items():
        ring = info.get("footprint")
        rec = by_plot.get(str(pid))
        if not ring or rec is None:
            missing += 1
            continue
        feats.append({"type": "Feature",
                      "properties": {"plot": int(pid), "row": rec["_row"], "col": rec["_col"],
                                     "accession": rec["_accession"]},
                      "geometry": {"type": "Polygon", "coordinates": [ring]}})
    if not feats:
        raise ValueError("no stitched plot matched a row of the design (do the plot numbers agree?)")
    directory = "Processed/" + "/".join(str(scope[k]) for k in ("year", "experiment", "location", "population", "date",
                                                                "platform", "sensor")) + "/"
    v = c.save_version(directory, {"boundaries": {"type": "FeatureCollection", "features": feats},
                                   "grid": {"rows": len(design["tiers"]), "cols": len(design["beds"]), "spacing_m": 0,
                                            "angle_deg": 0, "row_offset": 0, "col_offset": 0,
                                            "fill_pattern": "row-major"},
                                   "created_from": "grid",
                                   "field_design": {"source": Path(args["design_path"]).name, "by": "agent"}},
                       args.get("name") or f"agent: {len(feats)} stitched plots")
    c.post_json("/api/plot_geometry/versions/activate", {"directory": directory, "version": v["version"]})

    def update(pr):
        run = pr.ensure_run(scope, "ground")
        pr.set_step(run, "plot_boundary_prep", "completed",
                    outputs={"activeVersion": v["version"], "activeVersionName": v.get("name")})
        return pr.describe(run)

    return {"result": {"plots": len(feats), "not_matched": missing, "version": v["version"], "directory": directory,
                       "process_tab": process_tab(update)}}


def associate_plots(args: dict) -> dict:
    """ASSOCIATE_BOUNDARIES: file each stitched plot mosaic under its plot
    (PlotImages/plot_{n}_accession_{a}.png), using the active boundaries."""
    prefix = args["stitch_prefix"].rstrip("/") + "/"
    scope = scope_from_path(prefix)
    if scope is None:
        raise ValueError("stitch_prefix must be Processed/…/AgRowStitch_v{N}/")
    c = client()
    directory = "Processed/" + "/".join(str(scope[k]) for k in ("year", "experiment", "location", "population", "date",
                                                                "platform", "sensor")) + "/"
    v = c.load_version(directory)
    boundaries = (v.get("state_snapshot") or {}).get("boundaries") or {}
    if not boundaries.get("features"):
        raise ValueError("no active plot boundaries: run ground_plot_boundaries first")
    exp = next((e for e in c.entities("experiments") if e.get("experiment_name") == scope["experiment"]), None)
    params = {"stitch_prefix": prefix, "boundaries": boundaries, "plot_images_prefix": directory + "PlotImages/",
              "boundary_version": v.get("version")}
    job = c.submit_job("ASSOCIATE_BOUNDARIES", params, str(exp["id"]) if exp else None)

    def update(pr):
        return pr.record_job(scope, job, kind="ground")

    return {"result": {"job_id": job.get("id"), "status": job.get("status"), "plots": len(boundaries["features"]),
                       "output": directory + "PlotImages/", "process_tab": process_tab(update)}}


_tool(mock_field_design, "Write a stand-in field design (CSV) for a rover track that has none: bed = pass, tier = "
      "position in the bed, placeholder genotypes. For trying the pipeline only.",
      {"genotypes": _int, "save_as": _str}, [], track=True, group="ground")
_tool(ground_plot_boundaries, "Rover: make plot boundaries from a stitch run's georeferenced plot footprints, labelled "
      "from a field design; saves and activates them (the ground run's Plot Boundary Prep).",
      {"stitch_prefix": _str, "design_path": _str, "name": _str}, ["stitch_prefix", "design_path"],
      track=False, group="ground")
_tool(associate_plots, "Rover: file each stitched mosaic under its plot using the active boundaries "
      "(ASSOCIATE_BOUNDARIES; needs the stitch worker). Output: …/PlotImages/.",
      {"stitch_prefix": _str}, ["stitch_prefix"], track=False, group="ground")
