"""field_layout reads the shapes breeders keep: plot tables with any column
names, drawn maps (tiers down / beds across, with a crop label row; beds down
with plot + variety pairs), several files joined by plot number, seasons kept
apart, and files that disagree. Synthetic files, built here."""
import csv
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

from gemi_agent.field_layout import read_layout, summarize  # noqa: E402

tmp = Path(tempfile.mkdtemp(prefix="gemi-layout-test-"))


def xlsx(name, rows, merges=(), sheets=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rs, ms in sheets or [("Sheet1", rows, merges)]:
        ws = wb.create_sheet(title)
        for r in rs:
            ws.append(r)
        for m in ms:
            ws.merge_cells(m)
    p = tmp / name
    wb.save(p)
    return str(p)


# 1. A drawn map: notes, a crop label row merged over its beds, sizes, "Tier/bed #", plots; B = border.
grid = [["", "Lab plantings, Summer 2023"], ["", "B = Border"], [],
        ["", "", "CP B", "COWPEA MAGIC", "", "SORGHUM", ""],
        ["Tier/bed dim. (ft)", "", 5, 5, 5, 5, 5, "Total width: 25'"],
        ["", "Tier/bed #", 1, 2, 3, 4, 5],
        [12.5, 1, "B", "B", "B", "B", "B"],
        [15, 2, "B", 101, 104, 3001, 3004],
        [15, 3, "B", 102, 103, 3002, 3003]]
m = xlsx("map.xlsx", grid, ["D4:E4", "F4:G4"])
lay = read_layout(m)
plots = {p["plot"]: p for p in lay["plots"]}
assert (plots["101"]["tier"], plots["101"]["bed"], plots["101"]["crop"]) == (2, 2, "Cowpea"), plots["101"]
assert (plots["3003"]["tier"], plots["3003"]["bed"], plots["3003"]["crop"]) == (3, 5, "Sorghum"), plots["3003"]
assert len([p for p in lay["plots"] if not p.get("border")]) == 8

# 2. A plot list with other column names and a combined "Row x-bed y" column, joined with the map.
t = tmp / "entries.csv"
with open(t, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["plots", "entry.list.vect", "Row.bed.number"])
    w.writerow([101, "MAGIC001", "Row 2-bed 2"])
    w.writerow([103, "MAGIC077", ""])
lay = read_layout([m, str(t)])
plots = {p["plot"]: p for p in lay["plots"]}
assert plots["101"]["genotype"] == "MAGIC001" and plots["103"]["genotype"] == "MAGIC077"
assert (plots["103"]["tier"], plots["103"]["bed"]) == (3, 3)          # position from the map
assert not lay["conflicts"]

# 3. Beds down, a (plot, variety) pair per tier, drawn one column right of its header.
pairs = [["Field 57"], ["Bed", "Tier 1", "", "Tier 2", ""],
         [107, "Buffer", "", "Buffer", ""],
         [109, "Buffer", 1001, "MAGIC179", 1044, "MAGIC322"],
         [110, "Buffer", 1002, "MAGIC147", 1043, "MAGIC336"]]
lay = read_layout(xlsx("pairs.xlsx", pairs))
plots = {p["plot"]: p for p in lay["plots"]}
assert (plots["1044"]["tier"], plots["1044"]["bed"], plots["1044"]["genotype"]) == (2, 109, "MAGIC322"), plots["1044"]

# 4. Several seasons and sites in one table: plot numbers repeat and stay apart; a filter keeps one.
master = xlsx("master.xlsx", [["Year", "Location", "Field.plot.number", "Bed", "Tier", "Crop", "Genotype"],
                              [2022, "Davis", 1, 4, 2, "Cowpea", "A"], [2023, "Davis", 1, 6, 3, "Cowpea", "B"],
                              [2023, "Kearney", 1, 9, 9, "Cowpea", "C"]])
assert len(read_layout(master)["plots"]) == 3
one = read_layout(master, year="2023", location="Davis")["plots"]
assert len(one) == 1 and (one[0]["bed"], one[0]["genotype"]) == (6, "B")
s = summarize(read_layout(master))
assert set(s["environments"]) == {"Davis 2022", "Davis 2023", "Kearney 2023"}

# 5. Two files with bed and tier the other way round: conflicts, and the swap is named.
a = xlsx("a.xlsx", [["Plot", "Bed", "Tier"]] + [[n, 60 + n, 10 + n] for n in range(1, 6)])
b = xlsx("b.xlsx", [["Plot", "Bed", "Tier"]] + [[n, 10 + n, 60 + n] for n in range(1, 6)])
s = summarize(read_layout([a, b]))
assert s["position_conflicts"] == 5 and s["bed_tier_swapped"], s

# 6. A workbook with a sheet per group: the sheet name labels its plots.
wb = xlsx("groups.xlsx", None, sheets=[("Sorghum", [["Plot", "Bed", "Tier"], [1, 1, 1]], []),
                                       ("Check", [["Plot", "Bed", "Tier"], [2, 1, 2]], [])])
crops = {p["plot"]: p["crop"] for p in read_layout(wb)["plots"]}
assert crops == {"1": "Sorghum", "2": "Check"}, crops

print("field layout ok")
